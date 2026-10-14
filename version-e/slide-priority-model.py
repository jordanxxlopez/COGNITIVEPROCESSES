"""Guide-based study ranking. Scores are review priorities, not exam probabilities."""
import argparse,html,json,math,re
from collections import Counter
from pathlib import Path
from docx import Document
IDS=(3,5,2,7,4,8,1,0)
GUIDE_ANCHORS=[(3, [10], 'What cognitive psychology studies'), (3, [12, 16], 'Donders:'), (3, [18], 'Wundt and structuralism'), (3, [19, 20], 'Ebbinghaus:'), (3, [21], 'William James'), (3, [22], "Watson's critique"), (3, [25, 26, 27], 'Classical conditioning'), (3, [29, 31], 'Tolman and cognitive maps'), (3, [33], "Chomsky's critique of Skinner"), (3, [34], 'The cognitive revolution'), (3, [35], 'The information-processing approach and the computer analogy'), (3, [36, 37], 'Dichotic listening and Broadbent'), (3, [38, 39], 'Early artificial intelligence'), (3, [41], 'Structural vs. process models'), (5, [9, 10, 11], 'The inverse projection problem;'), (5, [13, 14, 15], 'Bottom-up vs. top-down; direct'), (5, [16], 'Recognition-by-components'), (5, [18, 19, 21], 'Top-down examples:'), (5, [23], 'Helmholtz:'), (2, [2], 'Gestalt organization:'), (2, [3, 4, 5, 8, 9], 'Regularities in the environment:'), (2, [11, 12], 'Bayesian inference:'), (2, [14], 'How the four approaches compare'), (2, [16, 17, 18], 'Cross-cultural illusion findings;'), (2, [20, 21], 'Experience-dependent plasticity'), (2, [22, 23, 25], 'Face- and place-selective areas;'), (2, [28, 30], 'Movement and manual exploration'), (2, [31], 'Perception and attention:'), (5, [45, 46], '“What” and “where/how” pathways;'), (5, [46], 'Mirror neurons,'), (7, [2], 'Types of attention:'), (7, [4, 5, 12], 'Dichotic listening, shadowing,'), (7, [8, 13, 16, 19, 21], 'Early, intermediate, and late selection models'), (7, [22], 'How the early–late debate was resolved'), (7, [25, 26], 'Load theory:'), (7, [24, 38, 39], 'Automatic vs. controlled processing;'), (7, [27], 'Overt vs. covert attention;'), (7, [28, 29, 31], 'Bottom-up salience vs. top-down guidance'), (7, [33, 34], 'Precueing; object-based attention'), (4, [5, 7], 'Divided attention in the real world:'), (4, [8, 9], 'Inattentional blindness vs. change blindness'), (4, [11, 12, 13, 14, 20, 21, 23], 'The binding problem;'), (4, [25, 29, 30, 31], 'Physiology of attention:'), (8, [2], 'The questions psycholinguistics asks:'), (8, [3, 4, 5], 'Phonemes and morphemes;'), (8, [7], 'Bottom-up and top-down processes in recognizing words'), (8, [6, 8], 'Phonemic restoration;'), (8, [11, 13], 'Corpora, word frequency,'), (8, [17, 18, 19], 'Lexical ambiguity:'), (8, [20], 'Syntax vs. semantics'), (8, [21, 23], 'Neuropsychological and ERP evidence'), (8, [24], 'Parsing: garden-path sentences'), (1, [2, 4, 5, 6, 7], 'Syntax-first vs. interactionist approaches;'), (1, [8, 9], 'Prediction and experience in sentence processing'), (1, [10, 11], 'Inference; situation models; simulation'), (1, [13, 14], 'Brain evidence that readers build models'), (1, [15, 16, 18], 'The given–new contract;'), (1, [19, 20, 22, 23], 'The Sapir-Whorf hypothesis:'), (1, [24, 25, 26, 28], "Chomsky's critique; Marr's levels"), (0, [12], 'Selective listening in a multi-talker'), (0, [15], 'Media multitasking and attention across three samples'), (0, [16], 'Whether English and Mandarin speakers'), (0, [14], "Children's pseudohomophone reading"), (0, [13], 'Selective attention in emergency wayfinding')]
STOP=set('the a an of and to in is are for with from this which what how as be by at only it on that our when not or its have can than those these into one their more some'.split())
def tokens(s):return Counter(w for w in re.findall(r'[a-z]+',html.unescape(re.sub('<[^>]*>',' ',s)).lower()) if w not in STOP)
def cosine(a,b):return sum(v*b.get(k,0) for k,v in a.items())/(math.sqrt(sum(v*v for v in a.values())*sum(v*v for v in b.values())) or 1)
def rank(data,guide):
 previously_marked={(m['id'],s['page']) for m in data['modules'] for s in m['sections'] if s.get('examPriorityE',{}).get('mustKnow')}
 guide_text=[p.text.strip() for p in Document(guide).paragraphs if p.text.strip()]
 blocks={k:[] for k in ('themes','history','perception','attention','language','articles')};group=None
 for p in Document(guide).paragraphs:
  t=p.text.strip()
  if t.startswith('I. Big'):group='themes'
  elif t.startswith('II. Chapter'):group='history'
  elif t.startswith('III. Chapter'):group='perception'
  elif t.startswith('IV. Chapter'):group='attention'
  elif t.startswith('V. Chapter'):group='language'
  elif t.startswith('VI. Assigned'):group='articles'
  elif group and t:blocks[group].append(t)
 groups={3:'history',5:'perception',2:'perception',7:'attention',4:'attention',8:'language',1:'language',0:'articles'}
 report={'method':'Guide and quiz ranking: 60% best study-guide similarity, 25% direct quiz coverage weighted by existing question priority, 15% direct essentials reference. Preserve previous red markers and the lecture top-40% baseline; add reviewed slide anchors for every substantive study-guide topic, all five assigned articles, and non-memory labs represented in the question bank.', 'interpretation':'Must know identifies core review priorities. Scores are not probabilities of appearing on the exam.', 'validated_prediction':False,'source_material':'Supplied Exam 1 study guide, lecture scripts, 208-question bank, essentials guide, and existing guide-weighted question priorities using six supplied example questions. No external exam dataset.','sections':[],'guide_coverage':[]}
 for m in data['modules']:
  if m['id'] not in IDS:continue
  candidates=[];guide_lines=blocks['themes']+blocks[groups[m['id']]]
  for s in m['sections']:
   vec=tokens(s['title']+' '+' '.join(s.get('versionE',s['bullets'])))
   matches=sorted(((cosine(vec,tokens(t)),t) for t in guide_lines),reverse=True)
   qs=[q for q in data['questions'] if q['module']==m['id'] and q['page']==s['page']]
   qsignal=min(1,sum(q.get('priority',1) for q in qs)/4)
   refs=[e['id'] for e in data['essentials'] if e['source']==f"{m['id']}:{s['page']}"]
   score=.6*matches[0][0]+.25*qsignal+.15*bool(refs)
   record={'module':m['id'],'page':s['page'],'title':s['title'],'score':round(score,6),'guideTopic':matches[0][1],'questionIds':[q['id'] for q in qs],'essentials':refs,'mustKnow':False}
   candidates.append((score,s,record))
  count=math.ceil(len(candidates)*.4)
  for i,(_,s,record) in enumerate(sorted(candidates,key=lambda x:(-x[0],x[2]['page']))):
   prior=(m['id'],s['page']) in previously_marked
   baseline=m['id']!=0 and i<count
   lab=m['id']==0 and bool(record['questionIds']) and s['page'] in (1,2,3,4,5,6,7,8,9,11)
   record['mustKnow']=prior or baseline or lab
   record['reasons']=[label for label,active in [('Existing core-review marker',prior),('Guide/quiz-weighted lecture ranking',baseline),('Non-memory lab in the practice-question bank',lab)] if active]
  report['sections'].extend(x[2] for x in candidates)
 by_key={(s['module'],s['page']):s for s in report['sections']}
 for mid,pages,prefix in GUIDE_ANCHORS:
  topic=next(t for t in guide_text if t.startswith(prefix))
  report['guide_coverage'].append({'guideTopic':topic,'module':mid,'pages':pages})
  for page in pages:
   item=by_key[(mid,page)];item['mustKnow']=True
   item.setdefault('guideTopics',[]).append(topic)
   item['guideTopic']=topic
   if 'Explicit study-guide coverage' not in item['reasons']:item['reasons'].append('Explicit study-guide coverage')
 for m in data['modules']:
  for s in m['sections']:
   if (m['id'],s['page']) in by_key:
    record=by_key[(m['id'],s['page'])]
    s['examPriorityE']={k:record[k] for k in ('mustKnow','score','guideTopic','questionIds','essentials','reasons')}
 return report
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--data',required=True);p.add_argument('--guide',required=True);p.add_argument('--output',required=True);a=p.parse_args()
 d=json.loads(Path(a.data).read_text());r=rank(d,a.guide);Path(a.output).write_text(json.dumps(r,ensure_ascii=False,indent=2));print(sum(x['mustKnow'] for x in r['sections']),'core review slides')
